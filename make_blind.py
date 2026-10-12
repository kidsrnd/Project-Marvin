"""
make_blind.py — 교사 블라인드 평가지 만들기 (교수님 4주차 피드백 ②④)
A / B1 / B1R(B1 별도검수 후) / B2 문항을 섞고 조건을 숨긴다.
  blind_sheet.xlsx  ← 선생님이 채점하는 파일 (조건 정보 없음)
  blind_key.csv     ← 번호별 조건 (채점 끝날 때까지 열지 않기)
실행: python make_blind.py outputs/20261005_1336
"""
import json, random, csv, sys
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

COLS = [("번호",6),("유형",7),("문제",40),("선지",44),("표시된 정답",30),
        ("정답이 맞나 (O/X/?)",10),
        ("[치명] 본문 근거 없는 정답",11),("[치명] 정답이 2개 이상",11),
        ("[중대] 본문만으로 구별 안 되는 오답",12),("[중대] 문제가 답을 알려줌",11),("[중대] 단어 찾기로 풀림",11),
        ("[경미] 본문에 없는 오답(규칙 위반)",12),("[경미] 어휘·길이 문제",10),
        ("서사 요소 (인물/배경/사건/감정/인과/결말/예측)",14),("명시/암시",9),("메모",30)]

def load(run):
    pool = []
    for cond, pat in [("A","A_{}.json"),("B1","B1_{}.json"),("B1R","review/B1R_{}.json"),("B2","B2_{}.json")]:
        for r in (1,2,3):
            for it in json.loads((run/pat.format(r)).read_text(encoding="utf-8"))["items"]:
                pool.append((cond, r, it))
    return pool

def answer_text(it):
    if it["type"] == "mc": return it.get("answer","")
    if "sample_answer" in it: return it["sample_answer"]
    if "rubric" in it: return "채점 핵심: " + it["rubric"].get("key_elements","")
    return it.get("answer","")

def main(run_dir):
    run = Path(run_dir); out = run/"blind"; out.mkdir(exist_ok=True)
    pool = load(run); random.Random(20261012).shuffle(pool)
    wb = Workbook(); ws = wb.active; ws.title = "블라인드 평가"
    f11 = Font(name="맑은 고딕", size=11); fb = Font(name="맑은 고딕", size=11, bold=True)
    for c,(h,w) in enumerate(COLS,1):
        cell = ws.cell(1,c,h); cell.font=fb; cell.alignment=Alignment(wrap_text=True,vertical="center")
        cell.fill = PatternFill("solid", fgColor="FDE2E2" if "치명" in h else "FFF2CC" if "중대" in h else "E2EFDA" if "경미" in h else "DDEBF7")
        ws.column_dimensions[ws.cell(1,c).column_letter].width = w
    ws.row_dimensions[1].height = 60
    oxq = DataValidation(type="list", formula1='"O,X,?"', allow_blank=True)
    yn  = DataValidation(type="list", formula1='"있음"', allow_blank=True)
    nar = DataValidation(type="list", formula1='"인물,배경,사건,감정,인과,결말,예측"', allow_blank=True)
    ei  = DataValidation(type="list", formula1='"명시,암시"', allow_blank=True)
    for dv in (oxq,yn,nar,ei): ws.add_data_validation(dv)
    key = []
    for n,(cond,r,it) in enumerate(pool,1):
        row = n+1
        vals = [n, {"mc":"객관식","short":"단답","cr":"서술형"}[it["type"]], it["question"],
                "\n".join(it.get("options",[])), answer_text(it)]
        for c,v in enumerate(vals,1):
            cell = ws.cell(row,c,v); cell.font=f11; cell.alignment=Alignment(wrap_text=True,vertical="top")
        for c in range(6, len(COLS)+1):
            ws.cell(row,c).font=f11; ws.cell(row,c).alignment=Alignment(wrap_text=True,vertical="top")
        oxq.add(f"F{row}"); nar.add(f"N{row}"); ei.add(f"O{row}")
        for col in "GHIJKLM": yn.add(f"{col}{row}")
        key.append([n, cond, r, it["id"], it["type"]])
    ws.freeze_panes = "C2"
    g = wb.create_sheet("채점 방법")
    for i,line in enumerate([
        "1. 조건(A/B1/B2 등)은 일부러 숨겼습니다. 어느 로봇이 만든 문제인지 추측하지 말고 문제만 보세요.",
        "2. '정답이 맞나'에 O / X / ? 를 고르세요.",
        "3. 문제가 있는 칸에만 '있음'을 고르세요. 없으면 비워 두면 됩니다.",
        "   [치명] 하나라도 있으면 그 문항은 탈락입니다.",
        "4. 서사 요소와 명시/암시는 문제마다 하나씩 고르세요.",
        "   명시 = 답이 본문에 그대로 있음 / 암시 = 읽고 추론해야 함",
        "5. 다 끝나면 파일을 저장하고 Claude에게 '채점 끝났어'라고 말해 주세요.",
        "6. blind_key.csv 는 채점이 끝날 때까지 열지 마세요."],1):
        g.cell(i,1,line).font=f11
    g.column_dimensions["A"].width = 100
    wb.save(out/"blind_sheet.xlsx")
    with open(out/"blind_key.csv","w",newline="",encoding="utf-8-sig") as f:
        w=csv.writer(f); w.writerow(["번호","조건","반복","원래문항번호","유형"]); w.writerows(key)
    print(f"{len(pool)}문항 → {out}")

if __name__ == "__main__":
    main(sys.argv[1])
